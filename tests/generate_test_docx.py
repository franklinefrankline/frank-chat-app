import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def generate_rich_docx():
    doc = docx.Document()

    # Title
    title = doc.add_heading('FRANK Think — Executive Project Brief', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Heading 1
    doc.add_heading('1. Overview & Strategic Goals', level=1)
    p1 = doc.add_paragraph('This document outlines the operational directives and performance benchmarks for the upcoming quarter. ')
    p1_run = p1.add_run('Key deliverables must adhere to enterprise quality standards and responsive UX specifications across desktop, tablet, and mobile devices.')
    p1_run.bold = True

    # Heading 2
    doc.add_heading('2. Key Initiatives & Milestones', level=2)
    doc.add_paragraph('The team has established the following milestone schedule for project execution:')

    # Numbered List
    doc.add_paragraph('Phase 1: Architecture review and UI/UX modernization.', style='List Number')
    doc.add_paragraph('Phase 2: In-chat document rendering engine integration (PDF, Word, Excel, PPTX).', style='List Number')
    doc.add_paragraph('Phase 3: Cross-platform responsiveness and performance optimization.', style='List Number')

    # Bullet List
    doc.add_heading('3. Functional Requirements', level=2)
    doc.add_paragraph('Zero horizontal page clipping on mobile viewports (320px to 430px).', style='List Bullet')
    doc.add_paragraph('Proper text wrapping, line-height, and padding preservation.', style='List Bullet')
    doc.add_paragraph('Responsive scaling for multi-column data tables.', style='List Bullet')
    doc.add_paragraph('Full preservation of formatting: bold, italic, lists, and headings.', style='List Bullet')

    # Mixed formatting paragraph with long URL
    doc.add_heading('4. Technical Reference & Links', level=2)
    p_link = doc.add_paragraph('For documentation, please refer to the internal repository and API specifications at ')
    p_link.add_run('https://github.com/frank-think-communications/enterprise-chat-application-platform/blob/main/docs/architecture-overview-specifications.md').italic = True
    p_link.add_run(' for full architectural diagrams.')

    # Wide Table
    doc.add_heading('5. Performance Benchmarks Table', level=2)
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Metric'
    hdr_cells[1].text = 'Target'
    hdr_cells[2].text = 'Mobile 375px'
    hdr_cells[3].text = 'Tablet 768px'
    hdr_cells[4].text = 'Desktop 1440px'

    data = [
        ('First Paint', '< 200ms', '140ms', '110ms', '85ms'),
        ('Document Load', '< 800ms', '620ms', '450ms', '320ms'),
        ('Page Fit', '100% visible', 'Pass (No Clip)', 'Pass', 'Pass'),
        ('Thumbnails', 'Generated', 'Pass', 'Pass', 'Pass'),
        ('Zoom 75%', 'Centered & Visible', 'Pass', 'Pass', 'Pass'),
    ]

    for metric, target, m375, t768, d1440 in data:
        row_cells = table.add_row().cells
        row_cells[0].text = metric
        row_cells[1].text = target
        row_cells[2].text = m375
        row_cells[3].text = t768
        row_cells[4].text = d1440

    # Page 2 (Break)
    doc.add_page_break()
    doc.add_heading('6. Sign-off & Appendix (Page 2)', level=1)
    p2 = doc.add_paragraph('This second page verifies multi-page pagination and thumbnail navigation in the FRANK Think document viewer.')
    p2.add_run(' All pages must render with independent page numbers and thumbnails.').bold = True

    output_path = 'tests/fixtures/rich_word_document.docx'
    doc.save(output_path)
    print(f'Successfully generated {output_path}')

if __name__ == '__main__':
    generate_rich_docx()
